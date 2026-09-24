#!/usr/bin/env python3
"""Structural checks on a built .docx: what should be in it is there, and no Markdown leaked through.

    python check.py out.docx [--md article.md] [--placeholders TEXT ...]

Prints ERROR / WARN / INFO lines; exits 1 when there is any ERROR. --md adds comparisons with the
source (formulas, pictures, tables, word count). --placeholders adds template leftovers to hunt for
(a sample author's name, 'XX'); a few generic ones are always checked. Cannot judge how it looks:
that is what render.py and your own eyes are for.
"""

import argparse
import re
import sys
import zipfile
from collections import Counter

from docx import Document

OMML = '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'
LEAKS = {
    'bold markers **': r'\*\*', 'backtick': r'`', 'display math $$': r'\$\$',
    'inline math $..$': r'(?<!\$)\$(?!\d)[^$\n]+\$', 'image syntax ![': r'!\[', 'link syntax ](': r'\]\(',
    'heading marker #': r'^\s*#{1,6}\s', 'table row |..|': r'^\s*\|.*\|\s*$',
    'unconverted list item': r'^\s*[-*+]\s+\w', 'LaTeX command': r'\\(?:frac|text|mathbb|mathbf|sum|cdot|tag|left|right)\b',
}
MISSING = re.compile(r'^\[(?:Figure|Рисунок|Диаграмма)')
GENERIC_PLACEHOLDERS = [r'\bx{2,}\b', r'\bхх\b', r'\bXX\b', r'Lorem ipsum', r'\(style:', r'\[N\]', r'TODO', r'\?\?\?']

findings = []


def report(level, message):
    findings.append(level)
    print(f'{level}: {message}')


def md_facts(path):
    text = open(path, encoding='utf-8').read()
    text = re.sub(r'\A---\n.*?\n---\n', '', text, flags=re.S)
    facts = {
        'display': len(re.findall(r'\$\$.+?\$\$', text, flags=re.S)),
        'inline': len(re.findall(r'(?<!\$)\$(?!\d)(?:[^$\\\n]|\\.)+?\$(?!\$)', re.sub(r'\$\$.+?\$\$', '', text, flags=re.S))),
        'images': len(re.findall(r'^!\[[^\]]*\]\([^)]+\)\s*$', text, flags=re.M)),
        'mermaid': len(re.findall(r'^```mermaid', text, flags=re.M)),
        'tables': len(re.findall(r'(?:^\|.*\n?)+', text, flags=re.M)),
    }
    plain = re.sub(r'```.*?```|\$\$.*?\$\$|!\[[^\]]*\]\([^)]*\)|[#*`|>_$\\]', ' ', text, flags=re.S)
    facts['words'] = len(re.findall(r'\w+', plain))
    return facts


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('docx')
    ap.add_argument('--md')
    ap.add_argument('--placeholders', nargs='*', default=[])
    args = ap.parse_args()

    doc = Document(args.docx)
    body = [(p.style.name, p.text) for p in doc.paragraphs]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    furniture = [p.text for s in doc.sections
                 for kind in ('header', 'first_page_header', 'even_page_header', 'footer')
                 for p in getattr(s, kind).paragraphs]
    texts = [t for _, t in body] + cells

    leaks = Counter()
    example = {}
    for style, t in [(s, t) for s, t in body] + [('table cell', c) for c in cells]:
        for name, pattern in LEAKS.items():
            if re.search(pattern, t):
                leaks[name] += 1
                example.setdefault(name, f'[{style}] {t.strip()[:80]}')
    for name, n in leaks.items():
        report('WARN', f'{n} place(s) look like leaked Markdown ({name}); e.g. {example[name]}')

    missing = [t for t in texts if MISSING.match(t.strip())]
    for t in missing:
        report('ERROR', f'a picture or diagram is missing: {t.strip()[:110]}')

    patterns = GENERIC_PLACEHOLDERS + [re.escape(p) for p in args.placeholders]
    for source, items in (('body', texts), ('header/footer', furniture)):
        for pat in patterns:
            hits = [t for t in items if re.search(pat, t)]
            if hits:
                report('WARN', f'template placeholder /{pat}/ left in {source} ({len(hits)}x); e.g. {hits[0].strip()[:80]}')

    default_style = doc.styles['Normal'].name
    plain = [t for s, t in body if s == default_style and t.strip()]
    if plain:
        report('WARN', f'{len(plain)} text paragraph(s) still in the default style "{default_style}"; e.g. {plain[0].strip()[:70]}')

    empty_run = longest = 0
    for _, t in body:
        empty_run = empty_run + 1 if not t.strip() else 0
        longest = max(longest, empty_run)
    if longest >= 3:
        report('WARN', f'{longest} empty paragraphs in a row')

    omml = zipfile.ZipFile(args.docx).read('word/document.xml').decode('utf-8').count('<m:oMath>')
    pictures = len(doc.inline_shapes)
    section = doc.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    for shape in doc.inline_shapes:
        if shape.width > text_width * 1.01:
            report('ERROR', f'a picture is wider than the text area ({shape.width.cm:.1f} > {text_width.cm:.1f} cm)')
    report('INFO', f'{len(body)} paragraphs, {len(doc.tables)} table(s), {pictures} picture(s), {omml} native formula(s)')

    if args.md:
        m = md_facts(args.md)
        formulas = m['display'] + m['inline']
        if formulas and omml == 0:
            report('ERROR', f'the source has {formulas} formula(s) but the document has no native equations '
                            f'(math engine fell back to plain text)')
        elif omml < formulas:
            report('WARN', f'the source has {formulas} formula(s), the document {omml} native equation(s)')
        expected = m['images'] + m['mermaid']
        if pictures < expected - len(missing):
            report('WARN', f'the source has {expected} picture/diagram block(s), the document {pictures} picture(s)')
        if len(doc.tables) != m['tables']:
            report('WARN', f'the source has {m["tables"]} table(s), the document {len(doc.tables)}')
        words = len(re.findall(r'\w+', ' '.join(texts)))
        if words < 0.9 * m['words']:
            report('WARN', f'the document has {words} words against {m["words"]} in the source: content may be dropped')

    usage = Counter(s for s, t in body if t.strip())
    report('INFO', 'styles by use: ' + ', '.join(f'{s} x{n}' for s, n in usage.most_common(12)))
    sys.exit(1 if 'ERROR' in findings else 0)


if __name__ == '__main__':
    main()
