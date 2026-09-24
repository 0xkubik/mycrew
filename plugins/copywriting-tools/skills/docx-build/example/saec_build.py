#!/usr/bin/env python3
"""A second worked example, not a tool: the build script for the SAEC 2026 conference template.

It differs from ispras_build.py on purpose: a Russian block and an English block on the title page,
spacer paragraphs copied from the template's sample body, abstract labels in bold italic, the
reference list arriving as an ordered Markdown list, and a loud stop for any block kind the
template has no rule for (this article has no tables; a silent drop would hide a missing rule).

    python saec_build.py --input article.md --template template.docx --output build/article.docx
"""

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

import docgen
from docgen import add_inline_runs, insert_formula

# ── This template's styles, read from its own sample body ──────────────────
S_UDK, S_AUTHOR, S_POSITION = 'UDK', 'AuthorName', 'Position'
S_TITLE, S_AFFIL, S_ABSTRACT = 'TITLE', 'Affiliation', 'Abstract_Keywords'
S_H1, S_H2, S_TEXT, S_EQ = 'Heading', 'Subheading', 'Text', 'Equation'
S_FIGURE_CAPTION, S_REF = 'Figure_Name', 'References'
ABSTRACT_LABELS = ('**Аннотация', '**Abstract', '**Ключевые', '**Keywords')


def blank(doc):
    doc.add_paragraph(style=S_TEXT)


def title_block(doc, blocks, i, last):
    """One language block of the title page. Returns the index just past its closing '---'."""
    while i < len(blocks) and blocks[i][0] != 'hr':
        kind, *rest = blocks[i]
        if kind == 'paragraph':
            text = rest[0]
            if text.startswith('УДК'):
                add_inline_runs(doc.add_paragraph(style=S_UDK), text)
            elif text.startswith(ABSTRACT_LABELS):
                text = re.sub(r'^\*\*(.+?)\*\*', r'***\1***', re.sub(r'  +', ' ', text.replace('\n', ' ')))
                add_inline_runs(doc.add_paragraph(style=S_ABSTRACT), text)
                if text.startswith(('***Ключевые', '***Keywords')) and not last:
                    blank(doc)
            elif text.lstrip().startswith('**'):                    # "**Name**," / "position;" pairs
                for line in text.split('\n'):
                    style = S_AUTHOR if line.lstrip().startswith('**') else S_POSITION
                    add_inline_runs(doc.add_paragraph(style=style), line.rstrip(),
                                    suppress_bold=True, suppress_italic=True)
                blank(doc)
            else:                                                   # affiliations
                for line in text.split('\n'):
                    add_inline_runs(doc.add_paragraph(style=S_AFFIL), line.rstrip())
                blank(doc)
        elif kind == 'heading' and rest[0] == 1:
            add_inline_runs(doc.add_paragraph(style=S_TITLE), rest[1])
            blank(doc)
        i += 1
    return i + 1


def body_block(doc, blocks, i, input_dir):
    kind, *rest = blocks[i]
    if kind in ('hr', 'empty'):
        return i + 1

    if kind == 'heading':
        level, title = rest
        if title == 'Список литературы':
            blank(doc)
            add_inline_runs(doc.add_paragraph(style=S_REF), f'**{title}**')
        else:
            add_inline_runs(doc.add_paragraph(style=S_H1 if level <= 2 else S_H2), title)
    elif kind == 'paragraph':
        add_inline_runs(doc.add_paragraph(style=S_TEXT), re.sub(r'  +', ' ', rest[0].replace('\n', ' ')))
    elif kind == 'math_block':
        p = doc.add_paragraph(style=S_EQ)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        insert_formula(p, rest[0].strip(), display=True)
    elif kind == 'image':
        alt, path = rest
        docgen.insert_image(doc, docgen.resolve_image_path(path, input_dir), alt or None,
                            caption_style=S_FIGURE_CAPTION, missing_style=S_TEXT,
                            max_width_cm=12, max_height_cm=9)
        blank(doc)
    elif kind == 'list':                                            # here the list is the references
        for n, item in enumerate(rest[1], 1):
            add_inline_runs(doc.add_paragraph(style=S_REF), f'{n}. {item}')
    else:
        sys.exit(f'block kind {kind!r} has no rule in this template yet: {str(blocks[i])[:80]}')
    return i + 1


def build(blocks, template, input_dir):
    doc = Document(template)
    docgen.clear_header_highlights(doc)
    docgen.replace_header_text(doc, first='')                       # the template's "Приложение 1" note
    docgen.clear_body(doc)

    i = 0
    for n in range(2):                                              # Russian block, English block
        i = title_block(doc, blocks, i, last=(n == 1))
    while i < len(blocks):
        i = body_block(doc, blocks, i, input_dir)
    return doc


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--input', '-i', required=True)
    ap.add_argument('--output', '-o', required=True)
    ap.add_argument('--template', '-t', required=True)
    args = ap.parse_args()

    print(f"math engine: {docgen.math_engine()}", file=sys.stderr)
    _, blocks = docgen.parse_md(args.input)
    doc = build(blocks, args.template, os.path.dirname(os.path.abspath(args.input)))
    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    doc.save(args.output)
    print(f"done: {len(blocks)} blocks -> {args.output}")


if __name__ == '__main__':
    main()
