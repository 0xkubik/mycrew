"""Page furniture and low-level Word XML: clearing the template body, headers, footers, page numbers."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

_HEADER_KINDS = ('header', 'first_page_header', 'even_page_header')


def clear_body(doc):
    """Drop the template's sample content but keep its styles, headers, footers and page setup."""
    body = doc.element.body
    for child in list(body):
        if child.tag.split('}')[-1] in ('p', 'tbl', 'sdt'):
            body.remove(child)


def set_even_odd_headers(doc):
    """Turn on separate headers and footers for even and odd pages."""
    settings = doc.settings.element
    if settings.find(qn('w:evenAndOddHeaders')) is None:
        settings.append(OxmlElement('w:evenAndOddHeaders'))


def clear_header_highlights(doc):
    """Templates often highlight their placeholder header text; remove the highlight."""
    for section in doc.sections:
        for kind in _HEADER_KINDS:
            try:
                for para in getattr(section, kind).paragraphs:
                    for run in para.runs:
                        rpr = run._r.get_or_add_rPr()
                        for hl in rpr.findall(qn('w:highlight')):
                            rpr.remove(hl)
            except Exception:
                pass


def replace_header_text(doc, *, default=None, even=None, first=None, size_pt=8):
    """Overwrite the placeholder text of the default / even-page / first-page header.

    Only paragraphs that already carry text are replaced; None leaves that header untouched.
    """
    wanted = {'header': default, 'even_page_header': even, 'first_page_header': first}
    for section in doc.sections:
        for kind, text in wanted.items():
            if text is None:
                continue
            try:
                for para in getattr(section, kind).paragraphs:
                    if para.text.strip():
                        para.clear()
                        para.add_run(text).font.size = Pt(size_pt)
            except Exception:
                pass


def add_field_code(run, field):
    """A Word field such as PAGE or NUMPAGES inside a run."""
    begin, instr, end = OxmlElement('w:fldChar'), OxmlElement('w:instrText'), OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr.text = field
    end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _wipe(part):
    body = part._element
    for child in list(body):
        if child.tag.split('}')[-1] in ('p', 'tbl', 'sdt', 'bookmarkStart', 'bookmarkEnd'):
            body.remove(child)
    body.append(OxmlElement('w:p'))


def add_page_numbers(doc, *, font='Times New Roman', size_pt=9, start=1,
                     align_default=WD_ALIGN_PARAGRAPH.RIGHT, align_even=WD_ALIGN_PARAGRAPH.LEFT,
                     align_first=WD_ALIGN_PARAGRAPH.RIGHT):
    """Replace the footers with a PAGE number: odd pages right, even pages left by default."""
    section = doc.sections[0]
    footers = [('footer', align_default), ('even_page_footer', align_even),
               ('first_page_footer', align_first)]
    for attr, align in footers:
        try:
            footer = getattr(section, attr)
            _wipe(footer)
            footer.is_linked_to_previous = False
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.clear()
            para.alignment = align
            run = para.add_run()
            run.font.name = font
            run.font.size = Pt(size_pt)
            add_field_code(run, 'PAGE')
        except Exception:
            pass

    sect_pr = section._sectPr
    pg = sect_pr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sect_pr.append(pg)
    pg.set(qn('w:start'), str(start))
