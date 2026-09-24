"""Pictures: path resolving, sizing into the page, Mermaid rendering. Template style names are arguments."""

import glob
import os
import shutil
import subprocess
import sys
import tempfile

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

from .inline import add_inline_runs


def resolve_image_path(img_path, input_dir):
    """Absolute paths as they are; relative ones from the Markdown's folder, then from the cwd."""
    if os.path.isabs(img_path):
        return img_path
    candidate = os.path.join(input_dir, img_path)
    if os.path.isfile(candidate):
        return os.path.abspath(candidate)
    if os.path.isfile(img_path):
        return os.path.abspath(img_path)
    return candidate                                     # missing; insert_image reports it


def insert_image(doc, path, caption=None, *, caption_style=None, missing_style=None,
                 max_width_cm=11, max_height_cm=10):
    """A centred picture, scaled into max_width x max_height, and an optional caption paragraph.

    A file that cannot be inserted leaves a visible '[Figure: ...]' line, never a silent gap.
    """
    try:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic = para.add_run().add_picture(path, width=Cm(max_width_cm))
        extent = pic._inline.find(qn('wp:extent'))
        if extent is not None:
            h, max_h = int(extent.get('cy', 0)), int(Cm(max_height_cm).emu)
            if h > max_h:
                extent.set('cx', str(int(int(extent.get('cx', 0)) * max_h / h)))
                extent.set('cy', str(max_h))
    except Exception as e:
        para = doc.add_paragraph(style=missing_style) if missing_style else doc.add_paragraph()
        para.add_run(f"[Figure: {path} - {e}]")

    if caption:
        cap = doc.add_paragraph(style=caption_style) if caption_style else doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_runs(cap, caption)


def find_mmdc():
    """The Mermaid CLI, from PATH or the usual nvm / Homebrew places; None when absent."""
    found = shutil.which('mmdc')
    if found:
        return found
    patterns = [os.path.expanduser('~/.nvm/versions/node/*/bin/mmdc'),
                '/usr/local/bin/mmdc', '/opt/homebrew/bin/mmdc']
    for pattern in patterns:
        for path in sorted(glob.glob(pattern), reverse=True):
            if os.access(path, os.X_OK):
                return path
    return None


def render_mermaid(code, width=1400):
    """Render Mermaid source to a PNG in a temp folder. Returns its path, or None with a warning."""
    mmdc = find_mmdc()
    if not mmdc:
        print("warning: mmdc not found, mermaid diagrams will be skipped", file=sys.stderr)
        return None
    tmp = tempfile.mkdtemp()
    src, out = os.path.join(tmp, 'diagram.mmd'), os.path.join(tmp, 'diagram.png')
    with open(src, 'w', encoding='utf-8') as f:
        f.write(code)
    height = '900' if 'sequenceDiagram' in code else '700' if ('flowchart' in code or 'graph' in code) else '600'
    try:
        r = subprocess.run([mmdc, '-i', src, '-o', out, '-b', 'white', '-w', str(width), '-H', height],
                           capture_output=True, timeout=60)
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"warning: mmdc failed: {e}", file=sys.stderr)
        return None
    if r.returncode == 0 and os.path.exists(out):
        return out
    print(f"warning: mmdc error: {r.stderr.decode(errors='replace')}", file=sys.stderr)
    return None
