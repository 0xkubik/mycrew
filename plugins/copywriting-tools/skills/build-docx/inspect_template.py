#!/usr/bin/env python3
"""Describe a Word template so a build script can be written against it.

    python inspect_template.py template.docx                   # overview
    python inspect_template.py template.docx --style NAME ...  # full detail for the named styles
    python inspect_template.py template.docx --all             # full detail for every custom style

Prints, as Markdown: page setup, header/footer text, the styles used by the sample body (with
resolved font, size, alignment, indent, spacing), the names of the other custom styles, and the
sample body's paragraphs with their style. Placeholder text shows in the headers and the sample.
"""

import argparse

from docx import Document
from docx.oxml.ns import qn


def resolved(style, getter):
    """First non-None value walking up the style's base chain."""
    while style is not None:
        value = getter(style)
        if value is not None:
            return value
        style = style.base_style
    return None


def cm(length):
    return None if length is None else round(length.cm, 2)


def pt(length):
    return None if length is None else round(length.pt, 1)


def style_line(s):
    props = {
        'font': lambda x: x.font.name, 'size': lambda x: pt(x.font.size),
        'bold': lambda x: x.font.bold, 'italic': lambda x: x.font.italic,
        'align': lambda x: x.paragraph_format.alignment,
        'first_indent_cm': lambda x: cm(x.paragraph_format.first_line_indent),
        'left_indent_cm': lambda x: cm(x.paragraph_format.left_indent),
        'before_pt': lambda x: pt(x.paragraph_format.space_before),
        'after_pt': lambda x: pt(x.paragraph_format.space_after),
        'line': lambda x: x.paragraph_format.line_spacing,
    }
    shown = ', '.join(f'{k}={v}' for k, g in props.items() if (v := resolved(s, g)) is not None)
    base = s.base_style.name if s.base_style is not None else '-'
    return f'- **{s.name}** (based on {base}): {shown or "inherits everything"}'


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('template')
    ap.add_argument('--style', nargs='+', metavar='NAME')
    ap.add_argument('--all', action='store_true')
    args = ap.parse_args()

    doc = Document(args.template)
    print(f'# Template: {args.template}\n')

    print('## Page setup')
    for n, s in enumerate(doc.sections, 1):
        print(f'- section {n}: {cm(s.page_width)} x {cm(s.page_height)} cm; margins top {cm(s.top_margin)}, '
              f'bottom {cm(s.bottom_margin)}, left {cm(s.left_margin)}, right {cm(s.right_margin)}; '
              f'different first page: {s.different_first_page_header_footer}')
    print(f"- separate even/odd headers: {doc.settings.element.find(qn('w:evenAndOddHeaders')) is not None}\n")

    print('## Headers and footers (text as the template has it)')
    for n, s in enumerate(doc.sections, 1):
        for kind in ('header', 'first_page_header', 'even_page_header',
                     'footer', 'first_page_footer', 'even_page_footer'):
            try:
                texts = [p.text for p in getattr(s, kind).paragraphs if p.text.strip()]
            except Exception:
                texts = []
            if texts:
                print(f'- section {n} {kind}: {texts}')
    print()

    paragraph_styles = [s for s in doc.styles if s.type == 1]
    used = {p.style.name for p in doc.paragraphs}
    if args.style:
        print('## Requested styles')
        by_name = {s.name: s for s in paragraph_styles}
        for name in args.style:
            print(style_line(by_name[name]) if name in by_name else f'- **{name}**: no such paragraph style')
    else:
        used_styles = [s for s in paragraph_styles if s.name in used]
        others = [s for s in paragraph_styles if s.name not in used and not s.builtin]
        print('## Styles used by the sample body')
        for s in used_styles:
            print(style_line(s))
        print(f'\n## Other custom paragraph styles ({len(others)}){" - detail" if args.all else " - names only"}')
        if args.all:
            for s in others:
                print(style_line(s))
        else:
            print(', '.join(s.name for s in others) or '(none)')
        print('\nAsk for detail with --style NAME ...  Pick styles by what they do, not by their names alone.')
    print()

    print('## Sample body: [style] text {explicit formatting of the first run}; empty paragraphs are kept')
    lines, empties = [], 0
    for p in doc.paragraphs:
        if not p.text.strip():
            empties += 1
            continue
        if empties:
            lines.append(f'- (empty paragraph x{empties})')
            empties = 0
        run = next((r for r in p.runs if r.text.strip()), None)
        fmt = ''
        if run is not None:
            explicit = {'bold': run.bold, 'italic': run.italic, 'size': pt(run.font.size), 'font': run.font.name}
            shown = ', '.join(f'{k}={v}' for k, v in explicit.items() if v is not None)
            fmt = f' {{{shown}}}' if shown else ''
        lines.append(f'- [{p.style.name}] {p.text.strip()[:90]}{fmt}')
    print('\n'.join(lines[:70]))
    if len(lines) > 70:
        print(f'- ... {len(lines) - 70} more lines')
    print(f'\nTables in body: {len(doc.tables)}; inline pictures: {len(doc.inline_shapes)}')


if __name__ == '__main__':
    main()
