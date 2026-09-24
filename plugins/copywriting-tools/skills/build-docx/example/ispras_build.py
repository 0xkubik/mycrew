#!/usr/bin/env python3
"""A worked example, not a tool: the build script for ONE template (the ISP RAS journal).

Everything template-specific lives here: style names, the title-page phases, bilingual captions,
header text, table width, reference-list rules. Everything generic comes from docgen. A new
template gets its own script of this shape, written from that template's own styles.

    python ispras_build.py --input article.md --template template.docx --output article.docx
"""

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

import docgen
from docgen import add_inline_runs, insert_formula

# ── This template's vocabulary ─────────────────────────────────────────────
S_TITLE, S_AUTHOR, S_ABSTRACT, S_TITLE_EN = 'ispHeader', 'ispAuthor', 'ispAnotation', 'ispHeader1'
S_H1, S_H2, S_H3 = 'ispSubHeader-1 level', 'ispSubHeader-2 level', 'ispSubHeader-3 level'
S_TEXT, S_CAPTION, S_CODE = 'ispText_main', 'ispPicture_sign', 'программа1'
S_REF, S_BULLET = 'ispLitList', 'ispList1'
TABLE_WIDTH_CM = 14
PRE_BODY = {"аннотация", "abstract", "ключевые слова", "keywords", "для цитирования", "for citation",
            "благодарности", "acknowledgements", "информация об авторах",
            "information about authors", "список литературы"}
FIG_RU, FIG_EN = re.compile(r'^\*?Рис\.\s+\d+'), re.compile(r'^\*?Fig\.\s+\d+')
TAB_RU = re.compile(r'^\*?Табл?\.\s+\d+')
DEFAULT_META = {
    'author_ru': 'Автор А.А.', 'title_short_ru': 'Заголовок статьи.',
    'journal_ru': 'Труды ИСП РАН, 2025, том XX, вып. X, с. хх–хх.',
    'author_en': 'Author A.A.', 'title_short_en': 'Article Title.',
    'journal_en': 'Trudy ISP RAN/Proc. ISP RAS, vol. XX, issue X, 2025. pp. xx-xx.',
    'journal_first_header': 'Труды ИСП РАН, том XX, вып. X, 2025 г. // Trudy ISP RAN/Proc. ISP RAS, vol. XX, issue X, 2025',
}


def bilingual_caption(text, ru, en):
    """'Рис. 5. Text' -> 'Рис. 5. Text / Fig. 5. Text.' (the English half is generated, not written)."""
    if re.search(rf' / {en}\.? ', text):
        return text.rstrip('.') + '.'
    m = re.match(rf'({ru}\.?\s+(\d+)\.?\s*)(.*)', text)
    if not m:
        return text
    rest = m.group(3).rstrip('.')
    return f'{ru}. {m.group(2)}. {rest} / {en}. {m.group(2)}. {rest}.'


def para(doc, text, style, bold_prefix=None, suppress_bold=False):
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
        add_inline_runs(p, text, force_not_bold=True, suppress_bold=suppress_bold)
    else:
        add_inline_runs(p, text, suppress_bold=suppress_bold)
    return p


def body_text(doc, text, indent_cm=0.7):
    """A body paragraph; a $$..$$ inside it becomes its own centred formula paragraph."""
    current = None
    for part in re.split(r'(\$\$[\s\S]+?\$\$)', text):
        m = re.match(r'^\$\$([\s\S]+?)\$\$$', part.strip())
        if m:
            current = None
            fp = doc.add_paragraph(style=S_TEXT)
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            insert_formula(fp, m.group(1).strip(), display=True)
        elif part:
            if current is None:
                current = doc.add_paragraph(style=S_TEXT)
                current.paragraph_format.first_line_indent = Cm(indent_cm)
                current.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_runs(current, part)


def abstract_para(doc, text, default_label):
    text = re.sub(r'  +', ' ', text.replace('\n', ' '))
    kw = re.match(r'^\*\*(.+?):\*\*\s*(.*)', text)
    if kw:
        para(doc, kw.group(2), S_ABSTRACT, bold_prefix=kw.group(1) + ': ')
    else:
        para(doc, text, S_ABSTRACT, bold_prefix=default_label)


def author_lines(doc, text):
    for line in text.split('\n'):
        if line.strip():
            add_inline_runs(doc.add_paragraph(style=S_AUTHOR), line.rstrip(), suppress_bold=True)


def build(meta, blocks, template, input_dir):
    meta = {**DEFAULT_META, **meta}
    doc = Document(template)
    docgen.clear_header_highlights(doc)
    docgen.replace_header_text(
        doc,
        default=f"{meta['author_ru']} {meta['title_short_ru']} {meta['journal_ru']}",
        even=f"{meta['author_en']} {meta['title_short_en']} {meta['journal_en']}",
        first=meta['journal_first_header'])
    docgen.set_even_odd_headers(doc)
    docgen.clear_body(doc)

    phase, mermaid_n, i = 'ru_title', 0, 0
    while i < len(blocks):
        kind, *rest = blocks[i]

        if phase == 'ru_title':
            if kind == 'heading' and rest[0] == 1:
                para(doc, rest[1], S_TITLE)
                phase = 'ru_authors'
        elif phase == 'ru_authors':
            if kind == 'hr':
                phase = 'ru_meta'
            elif kind == 'paragraph':
                author_lines(doc, rest[0])
        elif phase == 'ru_meta':
            if kind == 'hr':
                phase = 'en_title'
            elif kind == 'paragraph':
                abstract_para(doc, rest[0], 'Аннотация. ')
        elif phase == 'en_title':
            if kind == 'heading' and rest[0] == 1:
                para(doc, rest[1], S_TITLE_EN)
                phase = 'en_authors'
        elif phase == 'en_authors':
            if kind == 'heading' and rest[0] == 2 and rest[1].strip().lower() == 'abstract':
                phase = 'en_meta'
            elif kind == 'paragraph':
                author_lines(doc, rest[0])
        elif phase == 'en_meta':
            if kind == 'hr':
                phase = 'body'
            elif kind == 'heading' and rest[0] == 2 and re.match(r'^\d+\.', rest[1].strip()):
                phase = 'body'
                continue
            elif kind == 'paragraph':
                abstract_para(doc, rest[0], 'Abstract. ')
        else:
            i = body_block(doc, blocks, i, input_dir, mermaid_n)
            mermaid_n += kind == 'code_block' and rest[0] == 'mermaid'
            continue
        i += 1

    docgen.add_page_numbers(doc)
    return doc


def body_block(doc, blocks, i, input_dir, mermaid_n):
    """Emit blocks[i] in the body phase. Returns the index of the next block to process."""
    kind, *rest = blocks[i]

    if kind in ('hr', 'empty'):
        return i + 1

    if kind == 'heading':
        level, title = rest[0], re.sub(r'\s*\[[\dНн,\s\-–]+\]', '', rest[1]).strip()
        if level == 1:
            para(doc, title, S_TITLE)
        elif level == 2:
            plain = re.sub(r'^\d+\.\s*', '', title).strip()
            if plain.lower() in PRE_BODY:
                para(doc, plain, S_H1)
            else:
                for run in para(doc, title, S_H2).runs:
                    run.bold = True
        else:
            para(doc, title, S_H3)
        return i + 1

    if kind == 'code_block':
        lang, code = rest
        if lang == 'mermaid':
            png = docgen.render_mermaid(code)
            if png:
                docgen.insert_image(doc, png, max_width_cm=11, max_height_cm=7)
            else:
                doc.add_paragraph(style=S_CODE).add_run(f"[Диаграмма {mermaid_n + 1} — рендеринг недоступен]")
        else:
            doc.add_paragraph(style=S_CODE).add_run(code)
        return i + 1

    if kind == 'math_block':
        p = doc.add_paragraph(style=S_TEXT)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        insert_formula(p, rest[0].strip(), display=True)
        return i + 1

    if kind == 'blockquote':
        text = rest[0].strip()
        if text.startswith('*') and text.endswith('*'):
            text = text[1:-1]
        p = doc.add_paragraph(style=S_TEXT)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_runs(p, text, italic=True)
        return i + 1

    if kind == 'table':
        nxt = blocks[i + 1] if i + 1 < len(blocks) else None
        caption = nxt[1].strip('*').strip() if nxt and nxt[0] == 'paragraph' and TAB_RU.match(nxt[1]) else None
        if caption:
            p = doc.add_paragraph(style=S_CAPTION)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, bilingual_caption(caption, 'Табл', 'Table'))
        docgen.add_table(doc, docgen.table_rows(rest[0]), width_cm=TABLE_WIDTH_CM)
        return i + 2 if caption else i + 1

    if kind == 'image':
        alt, path = rest
        if alt and FIG_RU.match(alt):
            alt = bilingual_caption(alt, 'Рис', 'Fig')
        docgen.insert_image(doc, docgen.resolve_image_path(path, input_dir), alt or None,
                            caption_style=S_CAPTION, missing_style=S_TEXT)
        return i + 1

    if kind == 'list':
        for item in rest[1]:
            p = doc.add_paragraph(style=S_BULLET)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_runs(p, item)
        return i + 1

    text = rest[0]                                                   # kind == 'paragraph'
    if FIG_RU.match(text) or FIG_EN.match(text):
        p = doc.add_paragraph(style=S_CAPTION)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_runs(p, bilingual_caption(text.strip('*').strip(), 'Рис', 'Fig'))
    elif TAB_RU.match(text):
        p = doc.add_paragraph(style=S_CAPTION)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline_runs(p, text.strip('*').strip())
    elif text.startswith(('*\\*', '*\\\\*')):
        add_inline_runs(doc.add_paragraph(style=S_ABSTRACT), text.lstrip('*').rstrip('*').strip().lstrip('\\').strip())
    elif text.startswith(('*Версия:', '*Раздел')):
        add_inline_runs(doc.add_paragraph(style=S_ABSTRACT), text.strip('*').strip())
    elif re.match(r'^\[(?:Н?\d+|[A-Za-z]+\d*)\]\s', text):
        p = doc.add_paragraph(style=S_REF)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline_runs(p, re.sub(r'  +', ' ', text.replace('\n', ' ')))
    else:
        text = re.sub(r'  +', ' ', text.replace('\n', ' '))
        prev = next((blocks[j][0] for j in range(i - 1, -1, -1) if blocks[j][0] != 'empty'), None)
        after_formula = prev == 'math_block' and re.match(r'^[а-яё]', text)
        body_text(doc, text, indent_cm=0.0 if after_formula else 0.7)
    return i + 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--input', '-i', required=True)
    ap.add_argument('--output', '-o', required=True)
    ap.add_argument('--template', '-t', required=True)
    args = ap.parse_args()

    print(f"math engine: {docgen.math_engine()}", file=sys.stderr)
    meta, blocks = docgen.parse_md(args.input)
    doc = build(meta, blocks, args.template, os.path.dirname(os.path.abspath(args.input)))
    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    doc.save(args.output)
    print(f"done: {len(blocks)} blocks -> {args.output}")


if __name__ == '__main__':
    main()
