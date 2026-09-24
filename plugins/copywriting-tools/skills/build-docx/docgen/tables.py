"""Markdown pipe tables -> Word tables with explicit borders and a fixed total width."""

import re

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

from .inline import add_inline_runs

TWIPS_PER_CM = 567
_BORDERS = f'''<w:tblBorders {nsdecls('w')}>''' + ''.join(
    f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV')) + '</w:tblBorders>'


def table_rows(lines):
    """Cell text per row from pipe-table lines; the |---|---| separator row is dropped."""
    rows = []
    for line in lines:
        if re.match(r'^\|[\s\-:]+\|', line):
            continue
        rows.append([c.strip() for c in line.split('|')[1:-1]])
    return rows


def set_cell(cell, text, *, font, size_pt, bold=False):
    cell.text = ''
    para = cell.paragraphs[0]
    add_inline_runs(para, text, bold=bold)
    for run in para.runs:
        run.font.name = font
        run.font.size = Pt(size_pt)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)


def add_table(doc, rows, *, width_cm=14, font='Times New Roman', size_pt=9, header_bold=True):
    """Insert a centred, fully bordered table of equal-width columns; the first row is the header."""
    if not rows:
        return None
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    total = int(width_cm * TWIPS_PER_CM)
    props = table._tbl.tblPr if table._tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    width = OxmlElement('w:tblW')
    width.set(qn('w:w'), str(total))
    width.set(qn('w:type'), 'dxa')
    props.append(width)
    props.append(parse_xml(_BORDERS))

    col_w = max(1, total // cols)
    for column in table.columns:
        for cell in column.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.insert(0, tc_w)
            tc_w.set(qn('w:w'), str(col_w))
            tc_w.set(qn('w:type'), 'dxa')

    for r, row in enumerate(rows):
        for c, text in enumerate(row[:cols]):
            set_cell(table.cell(r, c), text, font=font, size_pt=size_pt, bold=header_bold and r == 0)
    return table
